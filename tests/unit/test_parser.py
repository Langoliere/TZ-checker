from docx import Document

from tz_checker.parsers.docx_parser import DocumentParseError, parse_docx


def test_parses_headings_and_table(tmp_path):
    path = tmp_path / "sample.docx"
    doc = Document(); doc.add_heading("1 Общие сведения", 1); doc.add_paragraph("Назначение изделия.")
    table = doc.add_table(rows=2, cols=2); table.cell(0, 0).text = "Параметр"; table.cell(0, 1).text = "Значение"; table.cell(1, 0).text = "Масса"; table.cell(1, 1).text = "10 кг"; doc.save(path)
    parsed = parse_docx(path, "id")
    assert parsed.sections[0].section_id == "1"
    assert parsed.tables[0].columns == ["Параметр", "Значение"]

def test_rejects_corrupt_file(tmp_path):
    path = tmp_path / "broken.docx"; path.write_bytes(b"broken")
    try: parse_docx(path, "id")
    except DocumentParseError: pass
    else: raise AssertionError("corrupt file accepted")
