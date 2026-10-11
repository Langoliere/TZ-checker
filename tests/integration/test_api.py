from io import BytesIO
from unittest.mock import patch

from docx import Document
from fastapi.testclient import TestClient

from tz_checker.api.app import app


def docx_bytes() -> bytes:
    stream = BytesIO()
    document = Document()
    document.add_heading("1 Общие сведения", 1)
    document.add_paragraph("Изделие должен соответствовать ГОСТ 1.0.")
    document.save(stream)
    return stream.getvalue()


def test_upload_and_list_documents():
    with patch("tz_checker.api.app.process_document_task.delay"), TestClient(app) as client:
        response = client.post(
            "/v1/documents",
            files={"file": ("test.docx", docx_bytes(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
        )
        assert response.status_code == 202
        assert response.json()["status"] == "uploaded"
        assert client.get("/v1/documents").status_code == 200


def test_rejects_wrong_extension():
    with TestClient(app) as client:
        response = client.post("/v1/documents", files={"file": ("test.txt", b"x", "text/plain")})
        assert response.status_code == 415
