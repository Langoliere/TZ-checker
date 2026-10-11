from collections import Counter
from io import BytesIO

from docx import Document as WordDocument

from tz_checker.models.db import Document, FindingRecord


# Markdown формируется без шаблонизатора, чтобы отчет оставался переносимым и проверяемым.
def markdown_report(document: Document, findings: list[FindingRecord]) -> str:
    counts = Counter(item.payload["severity"] for item in findings)
    lines = [f"# Отчет по проверке: {document.title or document.filename}", "", f"Всего замечаний: **{len(findings)}**", "", "## Сводка", ""]
    lines.extend(f"- {severity}: {counts[severity]}" for severity in ("critical", "major", "minor", "advisory"))
    lines += ["", "## Замечания", ""]
    for index, item in enumerate(findings, 1):
        f = item.payload
        lines += [f"### {index}. [{f['severity']}] {f['finding_type']}", f"- Раздел: {f.get('section_id') or 'Документ'}; location: `{f['location']}`", f"- Источник: {f['source']}", f"- Проблема: {f['issue']}", f"- Основание: {f['evidence']}", f"- Рекомендация: {f['recommendation']}", ""]
        if f.get("reference"):
            lines.append(f"- Эталон: {f['reference']['document_id']}: «{f['reference']['quote']}»")
    return "\n".join(lines)


def docx_report(document: Document, findings: list[FindingRecord]) -> bytes:
    report = WordDocument()
    report.add_heading(f"Отчет по проверке: {document.title or document.filename}", 0)
    report.add_paragraph(f"Всего замечаний: {len(findings)}")
    table = report.add_table(rows=1, cols=6)
    for cell, value in zip(table.rows[0].cells, ["Severity", "Раздел", "Location", "Проблема", "Источник", "Рекомендация"], strict=True):
        cell.text = value
    for record in findings:
        f = record.payload
        cells = table.add_row().cells
        values = [f["severity"], f.get("section_id") or "Документ", f["location"], f["issue"], f["source"], f["recommendation"]]
        for cell, value in zip(cells, values, strict=True):
            cell.text = value
    # Буфер позволяет вернуть DOCX напрямую через API без временного файла на диске.
    output = BytesIO()
    report.save(output)
    return output.getvalue()
